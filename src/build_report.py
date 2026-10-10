#!/usr/bin/env python3
"""Build the final DOCX report from analysis CSV outputs."""

from pathlib import Path
import csv
import json

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
FIG = ROOT / "figures"
REPORT = ROOT / "report"
REPORT.mkdir(exist_ok=True)


def read_csv(name):
    with (OUT / name).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, size=font_size)
        shade_cell(hdr[i], "D9EAF7")
    # Repeat header rows and prevent individual rows from splitting across pages.
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
        trPr = table.rows[-1]._tr.get_or_add_trPr()
        cantSplit = OxmlElement("w:cantSplit")
        trPr.append(cantSplit)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_picture(doc, path, caption, width=6.4):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True
        c.runs[0].font.size = Pt(9)


def main():
    rankings = read_csv("team_rankings.csv")
    metrics = {r["team_id"]: r for r in read_csv("team_metrics.csv")}
    compatibility = read_csv("compatibility_ranking_3470.csv")
    reliability = read_csv("reliability_risk.csv")
    overrated = read_csv("overrated_underrated.csv")
    recs = read_csv("recommendations_3470.csv")
    uncertainty = read_csv("ambiguity_uncertainty.csv")
    strategies = read_csv("match_strategy_examples.csv")
    quality = read_csv("data_quality_summary.csv")
    cleaning = read_csv("cleaning_log.csv")
    scout_bias = read_csv("scout_bias_summary.csv")
    archetypes = read_csv("team_archetypes.csv")
    with (OUT / "summary.json").open(encoding="utf-8") as f:
        summary = json.load(f)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(28)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.name = "Aptos Display"
    styles["Heading 2"].font.size = Pt(13)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.text = "DECODE Competition Intelligence Project | Team 3470"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.runs[0].font.size = Pt(8)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("DECODE Competition Intelligence Project")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Complete analysis of Candidate Dataset 1 for Team 3470")
    r.bold = True
    r.font.size = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Python-based, explainable analysis. No machine learning used.")
    doc.add_paragraph()

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "After cleaning 109 raw observations into 108 analysis-ready rows, the strongest overall robot is T08. "
        "Team 3470 ranks 6th overall in the balanced model and is strongest in cycle speed, reliability, and autonomous leave. "
        "Its clearest performance bottleneck is a late-event slowdown associated with the overflow/intake path, while its raw scoring-action volume is only mid-field despite fast cycles."
    )
    add_bullets(doc, [
        "Overall top five: T08, T14, T05, T17, T07.",
        "Best partners for 3470: T08, T14, T17, T02, T06.",
        "Most reliable: T02, T17, 3470, T08, T03.",
        "3470: 100% autonomous leave, 0 recorded breakdowns, 15.03 s average cycle time, 10.5 average recorded scoring actions, 50% motif completion.",
        "Primary 3470 priority: fix late-match throughput before chasing more speed.",
    ])

    doc.add_heading("1. Data Cleaning and Method", level=1)
    doc.add_paragraph(
        "The analysis starts with the supplied scouting export and preserves raw observations whenever the data does not justify a correction. "
        "One exact duplicate was removed. Missing values were not imputed. Suspicious values were flagged, and only one low-confidence extreme CL value was excluded from CL-based aggregates while its raw value remained in the cleaned dataset."
    )
    add_table(doc, ["Check", "Result", "Interpretation"], [[q["check"], q["result"], q["interpretation"]] for q in quality], [1.55, 1.25, 4.3], 8.5)

    doc.add_heading("Ranking model", level=2)
    doc.add_paragraph(
        "Because the supplied materials do not give official point values for CL, overflow, depot, motif, or leave actions, I did not invent point weights. "
        "Instead, each team metric is converted to a field percentile so different units can be combined transparently. Alliance score and win rate are excluded from robot strength because they describe the whole alliance."
    )
    add_table(doc, ["Component", "Weight", "What it uses"], [
        ["Scoring volume", "45%", "Mean field percentiles of CL, overflow, and depot actions"],
        ["Objectives", "15%", "Auto leave rate and motif completion rate"],
        ["Cycle speed", "15%", "Inverse percentile of average cycle time"],
        ["Reliability", "15%", "Breakdowns, driver consistency, penalties, action variability, cycle variability"],
        ["Defense", "10%", "Scout defense rating percentile"],
    ], [1.55, 0.8, 4.75], 9)
    doc.add_paragraph(
        "Ranking confidence is based on data completeness, average scout confidence, and rank stability under balanced, offense-heavy, and reliability-heavy weightings. "
        "This is a decision model, not an official FTC scoring model or a prediction of exact match points."
    )

    doc.add_heading("2. Data Quality Findings", level=1)
    key_cleaning = [
        "M23-T09-4 appeared twice as an identical row. One copy was removed.",
        "M14-T05-3 reports 28 CL actions while T05's other usable CL values are 6-7; the same row has the dataset's lowest scout confidence (2/5). Raw value retained; CL value excluded from the composite.",
        "T12 is missing two cycle times, and another T12 cycle time is 7.5 s, making any 'fastest robot' claim uncertain.",
        "Two very slow cycle-time outliers (T04 at 25.5 s and T07 at 25.6 s) were retained because they coincide with breakdown-prone robots and are plausible failure-performance observations.",
        "There are 17 missing qualitative notes, but no W/L records conflict with the listed alliance/opponent scores and team match numbers remain 1-6 for every team.",
        "Scout defense means range from 2.67 to 3.18, suggesting mild rater-calibration differences. Defense is therefore used as a supporting signal, not a precise standalone truth.",
        "All 36 reused M## prefixes have conflicting scorelines for observations with the same alliance color; 6 prefixes even have three observations labeled the same alliance. The prefix may be an arbitrary/reused label, not a global match ID. No rows were deleted for this ambiguity, but event scorelines need independent verification.",
    ]
    add_bullets(doc, key_cleaning)
    add_table(doc, ["Scout", "Obs.", "Avg defense", "Avg driver", "Avg confidence"], [
        [r["scout_id"], r["observations"], r["avg_defense_rating"], r["avg_driver_consistency"], r["avg_scout_confidence"]]
        for r in scout_bias
    ], [1.0, 0.7, 1.25, 1.25, 1.35], 9)

    doc.add_heading("3. Overall Team Ranking", level=1)
    add_picture(doc, FIG / "overall_ranking.png", "Figure 1. Overall evidence-weighted team ranking.")
    ranking_rows = []
    for r in rankings:
        ranking_rows.append([
            r["rank"], r["team_id"], r["overall_score"], r["avg_total_actions"], r["avg_cycle_sec"],
            r["reliability"], r["confidence"],
        ])
    add_table(doc, ["Rank", "Team", "Score", "Avg actions", "Cycle s", "Reliability", "Confidence"], ranking_rows, [0.55, 0.65, 0.7, 0.9, 0.75, 0.9, 0.9], 8.0)

    doc.add_heading("Why the top teams rank where they do", level=2)
    add_bullets(doc, [
        "T08 (#1, High confidence): 13.17 average actions, 4.68/5 defense, 83% motif completion, no breakdowns. Main risks are 19.73 s average cycle time, 2.2/5 driver consistency, and 5.67 average penalty points. It stays #1 in all three weight-sensitivity models.",
        "T14 (#2, Medium confidence): fast 15.53 s cycles, 6.17 CL, 4.33 depot, low 2.83 average penalties, no breakdowns. Its rank moves from 2nd to 6th under different weight priorities, so placement is less stable than T08.",
        "T05 (#3, Medium confidence): high scoring volume and no breakdowns, but 6.5 average penalties, 2.25/5 driver consistency, and the low-confidence 28-CL outlier reduce trust and fit.",
        "T17 (#4, Medium confidence): strong driver consistency (4.85/5), 12.0 actions, 83% motif, no breakdowns. Weak defense (1.25/5) and 6.0 average penalties are the main drawbacks.",
        "T07 (#5, Medium confidence): the highest recorded action volume among the top group (13.83) and a 17-action ceiling, but a 50% breakdown rate makes it the clearest high-risk/high-ceiling option.",
    ])

    doc.add_heading("4. Reliability, Risk, and Floor/Ceiling", level=1)
    top_rel = reliability[:8]
    add_table(doc, ["Rank", "Team", "Reliability", "Floor", "Ceiling", "Breakdown", "Avg penalty", "Risk tags"], [
        [r["reliability_rank"], r["team_id"], r["reliability_index"], r["action_floor_q25"], r["action_ceiling_max"], r["breakdown_rate"], r["avg_penalty"], r["risk_tags"].replace("high-floor", "high floor").replace("high-ceiling/high-risk", "ceiling risk").replace("penalty-risk", "penalty risk").replace("breakdown-risk", "breakdown risk")]
        for r in top_rel
    ], [0.5, 0.55, 0.8, 0.55, 0.55, 0.8, 0.75, 2.2], 7.5)
    add_bullets(doc, [
        "Most reliable: T02, T17, 3470. T02 has the best combined reliability index; T17 adds a high action floor; 3470 has zero breakdowns and the lowest action variability in the field.",
        "Best high-floor option: T08. Its 25th-percentile action floor is 13.0 with zero recorded breakdowns.",
        "High-ceiling/high-risk: T07 is the clearest case (17-action ceiling, 50% breakdown rate). T05 carries penalty/driver/data risk despite strong scoring. T06 has a 16-action ceiling but one breakdown and slower cycles.",
        "Major failure/penalty risk: T11 and T07 each have 50% breakdown rates; T04 combines a 33% breakdown rate with 8.5 average penalty points.",
    ])

    doc.add_heading("5. Overrated / Underrated Relative to Raw Match Results", level=1)
    doc.add_paragraph(
        "This does not refer to external reputation. It compares the individual-robot composite to a separate result-context rank built from win rate and average alliance score. "
        "A robot can look better or worse in match results because alliance partners and opponents are not controlled in this dataset. "
        "Additionally, repeated M## match prefixes have conflicting scorelines within one alliance color; the meaning of the prefix must be verified before treating these as official event results. Consequently, result-context over/underrating labels are provisional."
    )
    delta_rows = [r for r in overrated if abs(int(r["rank_difference"])) >= 3]
    add_table(doc, ["Team", "Robot rank", "Result rank", "Difference", "Interpretation"], [
        [r["team_id"], r["overall_robot_rank"], r["result_context_rank"], r["rank_difference"], r["interpretation"]]
        for r in delta_rows
    ], [0.75, 0.9, 0.9, 0.8, 3.4], 8.5)
    doc.add_paragraph(
        "The strongest underrated signal is T14 (#2 robot profile vs #13 result context). T02 and T17 also look better underneath than their match results suggest. "
        "T11 is the strongest overrated-by-results signal (#18 robot profile vs #8 result context), followed by T06 and T10."
    )

    doc.add_heading("6. Team Archetypes", level=1)
    add_table(doc, ["Team", "Archetype"], [[r["team_id"], r["archetype"]] for r in archetypes], [1.0, 5.5], 9)

    doc.add_heading("7. Team 3470 Performance Analysis", level=1)
    add_picture(doc, FIG / "team_3470_profile.png", "Figure 2. Team 3470's field-percentile profile.")
    m = metrics["3470"]
    doc.add_paragraph(
        f"Team 3470 ranks #{m['overall_rank']} overall. It averages {float(m['avg_total_actions']):.2f} recorded scoring actions and {float(m['avg_cycle']):.2f} s cycle time, "
        f"with {float(m['breakdown_rate']):.0%} recorded breakdowns, {float(m['auto_rate']):.0%} autonomous leave, {float(m['motif_rate']):.0%} motif completion, "
        f"and {float(m['avg_driver']):.2f}/5 driver consistency."
    )
    add_bullets(doc, [
        "Strongest areas: cycle speed (88th percentile), action consistency/reliability, 6/6 autonomous leave, and zero recorded breakdowns.",
        "Weakest areas: CL production (41st percentile), depot production (35th percentile), defense (41st percentile), and only 50% motif completion.",
        "Most important performance gap: 3470 is fast but only 45th percentile in scoring-volume index. The robot gets around the cycle quickly without converting that speed into top-tier action count.",
        "Late-event trend: cycle time rises and scoring actions drift down across the six matches; the same overflow-slowdown note appears in matches 4-6.",
    ])
    add_picture(doc, FIG / "team_3470_cycle_trend.png", "Figure 3. Team 3470 cycle time trends slower across the sample.")
    add_picture(doc, FIG / "team_3470_actions_trend.png", "Figure 4. Team 3470 recorded scoring actions drift down late in the sample.")

    doc.add_heading("8. Prioritized Team 3470 Recommendations", level=1)
    for r in recs:
        doc.add_heading(f"Priority {r['priority']}: {r['problem']}", level=2)
        add_bullets(doc, [
            f"Evidence: {r['supporting_evidence']}",
            f"Expected impact: {r['expected_competitive_impact']}",
            f"Action: {r['recommended_action']}",
            f"Confidence: {r['confidence']}",
        ])

    doc.add_heading("9. 3470 Alliance Compatibility Ranking and Picklist", level=1)
    add_picture(doc, FIG / "compatibility_ranking.png", "Figure 5. Highest-fit alliance partners for Team 3470.")
    comp_rows = []
    for r in compatibility:
        comp_rows.append([
            r["compatibility_rank"], r["team_id"], f"{float(r['compatibility_score']):.1f}", r["overall_rank"], r["fit_reason"], r["major_risk"], r["confidence"],
        ])
    add_table(doc, ["Pick", "Team", "Fit score", "Overall", "Why it fits", "Risk", "Conf."], comp_rows, [0.45, 0.55, 0.65, 0.55, 2.0, 1.65, 0.55], 7.0)
    doc.add_paragraph(
        "T08 is both the #1 overall robot and the #1 partner for 3470. That is not automatic: the partner model is separate and specifically rewards filling 3470's weaker CL/depot/defense areas. "
        "T08's high defense and scoring volume create the strongest complement, while T14 is a safer fast/low-penalty alternative and T17 brings elite driver consistency."
    )
    add_bullets(doc, [
        "Top tier: T08, T14, T17, T02.",
        "Strong second tier: T06, T05, T15, T07. T05 and T07 have more obvious risk; T06 is slower but high-volume; T15 is stable but penalty-prone.",
        "Avoid if better options remain: T13, T04, T11. T13 is useful defensively but low-volume; T04 combines low output, breakdowns, and penalties; T11 is last in overall strength and has a 50% breakdown rate.",
    ])

    doc.add_heading("10. Match Strategy Examples", level=1)
    for s in strategies:
        doc.add_heading(f"Matchup {s['matchup']}: {s['our_alliance']} vs {s['opposing_alliance']}", level=2)
        add_table(doc, ["Item", "Drive-team brief"], [
            ["Primary win condition", s["primary_win_condition"]],
            ["Roles", s["recommended_alliance_roles"]],
            ["Biggest threat", s["biggest_opponent_threat"]],
            ["Priority", s["strategic_priority"]],
            ["Major risk", s["major_risk"]],
            ["Backup plan", s["backup_plan"]],
            ["Confidence", s["confidence"]],
        ], [1.3, 5.65], 8.3)

    doc.add_heading("11. Ambiguity and Uncertainty", level=1)
    for i, u in enumerate(uncertainty, 1):
        doc.add_heading(f"Uncertainty {i}", level=2)
        add_table(doc, ["Field", "Assessment"], [
            ["Claim", u["claim"]],
            ["Confidence", u["confidence"]],
            ["Competing explanation", u["competing_explanation"]],
            ["Missing information", u["missing_information"]],
            ["Decision", u["decision"]],
        ], [1.45, 5.5], 8.5)

    doc.add_heading("12. Additional Data Needed", level=1)
    add_bullets(doc, [
        "Official point values or scoring-value mapping for each action type, if the goal is to estimate point contribution rather than action volume.",
        "Alliance partner identities and opponent strength for every match so alliance scores/wins can be adjusted for schedule context.",
        "Official global match IDs and score sheets to resolve the repeated M## prefix and same-alliance scoreline conflicts.",
        "Complete per-cycle timestamps and subsystem-failure timestamps, especially for T12 and Team 3470.",
        "Multiple scout observations of the same matches or a calibrated scoring rubric for defense and driver ratings.",
        "Match video for the T05 28-CL observation and Team 3470's repeated late-overflow slowdown.",
        "Phase-specific data (autonomous, teleop, endgame), defense faced, battery state, and jam/retry counts for stronger causal conclusions.",
    ])

    doc.add_heading("13. Final Decision Summary", level=1)
    add_bullets(doc, [
        "Strongest overall: T08 (High confidence within this dataset).",
        "Best alliance partner for 3470: T08; safest close alternatives are T14 and T02, while T17 adds high driver execution.",
        "3470's competitive identity: fast and mechanically reliable, but not yet translating its speed into top-tier scoring volume.",
        "Highest-priority fix: late-event overflow/intake throughput. Then improve CL/depot conversion and motif reliability.",
        "Use the picklist as a living decision aid, not a fixed truth. Six matches per team and subjective scouting ratings mean pre-match verification still matters.",
    ])

    doc.add_heading("Appendix A. AI Use and Explainability", level=1)
    doc.add_paragraph(
        "AI-assisted development was used for code scaffolding, analysis-method brainstorming, debugging, report drafting, and formatting. "
        "The final method is intentionally simple enough to audit: cleaning rules, percentile conversion, explicit weights, and rule-based confidence/archetype logic. "
        "Before submitting, the candidate should rerun the pipeline and complete the verification checklist in INTERVIEW_PREP.md so any first-person AI disclosure is accurate."
    )
    doc.add_paragraph(
        "Where this system could be wrong: the supplied dataset does not include official point weights, alliance partner identities, schedule strength, full cycle timing, or duplicate independent scout observations. "
        "The analysis also finds conflicting scorelines among records sharing M## prefixes; their meaning requires official verification. "
        "The analysis therefore avoids precise point forecasts and labels uncertain conclusions instead of forcing them."
    )

    out = REPORT / "DECODE_Competition_Intelligence_Report.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
